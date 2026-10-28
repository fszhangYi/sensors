#!/usr/bin/env python3
"""Build 构建指南.docx with formulas as PNG (no OMML).

Source of truth remains 构建指南.md (LaTeX). This script:
1. Renders \\(...\\) / \\[...\\] to guide_assets/formulas/*.png
2. Writes a temporary MD (display math → centered images; inline → placeholders)
3. Runs md_to-docx-report converter
4. Replaces placeholders with inline pictures; strips any leftover OMML
5. Writes 构建指南.docx (+ 兼容版 / 含公式修复版 copies)
"""

from __future__ import annotations

import hashlib
import re
import shutil
import sys
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Emu
from lxml import etree

ROOT = Path(__file__).resolve().parent
MD_SRC = ROOT / "构建指南.md"
FORMULA_DIR = ROOT / "guide_assets" / "formulas"
MD2DOCX = Path("/root/.cursor/skills/md-to-docx-report/scripts/md_to_docx.py")
PLACEHOLDER_FMT = "⟦EQ{0:03d}⟧"

DELIM_RE = re.compile(
    r"\$\$(.+?)\$\$|\\\[(.+?)\\\]|\\\((.+?)\\\)",
    re.S,
)

# LaTeX → matplotlib-friendly rewrite (no bmatrix / CJK in math font)
LATEX_REWRITE = {
    r"T_{\mathrm{base}} = \begin{bmatrix} R_{\mathrm{xml}} R_{\mathrm{install}} & 0 \\ 0 & 1 \end{bmatrix}":
        "__MATRIX_TBASE__",
    r"(\theta_{\mathrm{读}}-\mathrm{offset})\times\mathrm{dir}":
        r"(\theta_{\mathrm{read}}-\mathrm{offset})\times\mathrm{dir}",
}


def _slug(latex: str) -> str:
    h = hashlib.sha1(latex.encode("utf-8")).hexdigest()[:10]
    return f"eq_{h}"


def render_matrix_tbase(path: Path, dpi: int = 220) -> None:
    """Compose the 2×2 install matrix equation without LaTeX bmatrix."""
    fig, ax = plt.subplots(figsize=(7.2, 1.35))
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 3)
    ax.text(0.15, 1.5, r"$T_{\mathrm{base}}=$", fontsize=20, va="center", ha="left")
    ax.text(2.55, 1.5, "[", fontsize=42, va="center", ha="center", family="serif")
    ax.text(2.95, 2.15, r"$R_{\mathrm{xml}}\,R_{\mathrm{install}}$", fontsize=15, va="center", ha="left")
    ax.text(8.35, 2.15, r"$0$", fontsize=15, va="center", ha="center")
    ax.text(2.95, 0.85, r"$0$", fontsize=15, va="center", ha="left")
    ax.text(8.35, 0.85, r"$1$", fontsize=15, va="center", ha="center")
    ax.text(9.1, 1.5, "]", fontsize=42, va="center", ha="center", family="serif")
    fig.savefig(path, dpi=dpi, bbox_inches="tight", pad_inches=0.12, facecolor="white")
    plt.close(fig)


def render_latex(latex: str, path: Path, *, display: bool, dpi: int = 220) -> None:
    rewritten = LATEX_REWRITE.get(latex.strip(), latex.strip())
    if rewritten == "__MATRIX_TBASE__":
        render_matrix_tbase(path, dpi=dpi)
        return
    fontsize = 18 if display else 15
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.patch.set_facecolor("white")
    txt = fig.text(0.5, 0.5, f"${rewritten}$", fontsize=fontsize, ha="center", va="center", color="black")
    fig.canvas.draw()
    bbox = txt.get_window_extent(renderer=fig.canvas.get_renderer()).transformed(
        fig.dpi_scale_trans.inverted()
    )
    pad = 0.22 if display else 0.12
    fig.set_size_inches(max(bbox.width + pad, 0.4), max(bbox.height + pad, 0.25))
    fig.savefig(path, dpi=dpi, bbox_inches="tight", pad_inches=0.06, facecolor="white", edgecolor="none")
    plt.close(fig)


def collect_and_render(md_text: str) -> tuple[str, list[dict]]:
    """Return (md_with_placeholders_or_imgs, formula records)."""
    FORMULA_DIR.mkdir(parents=True, exist_ok=True)
    records: list[dict] = []
    parts: list[str] = []
    last = 0
    idx = 0
    for m in DELIM_RE.finditer(md_text):
        parts.append(md_text[last : m.start()])
        latex = next(g for g in m.groups() if g is not None).strip()
        display = m.group(0).startswith("\\[") or m.group(0).startswith("$$")
        slug = _slug(("D:" if display else "I:") + latex)
        png = FORMULA_DIR / f"{slug}.png"
        if not png.exists() or png.stat().st_size < 50:
            try:
                render_latex(latex, png, display=display)
            except Exception as e:
                # fallback: write error text image
                fig, ax = plt.subplots(figsize=(4, 0.6))
                ax.axis("off")
                fig.patch.set_facecolor("white")
                ax.text(0.01, 0.5, latex[:80], fontsize=10, va="center", family="monospace")
                fig.savefig(png, dpi=160, bbox_inches="tight", facecolor="white")
                plt.close(fig)
                print(f"WARN render failed ({e!r}), wrote fallback for {latex[:60]!r}", file=sys.stderr)
        rel = f"guide_assets/formulas/{png.name}"
        rec = {"index": idx, "latex": latex, "display": display, "path": png, "rel": rel}
        records.append(rec)
        if display:
            eq_no = sum(1 for r in records if r["display"]) + 1
            parts.append(
                "\n\n"
                f'<p align="center">\n'
                f'  <img src="{rel}" alt="式 ({eq_no})" width="70%"/>\n'
                f"</p>\n"
                f'<p align="center"><em>式 ({eq_no})</em></p>\n\n'
            )
        else:
            parts.append(PLACEHOLDER_FMT.format(idx))
        idx += 1
        last = m.end()
    parts.append(md_text[last:])
    return "".join(parts), records


def strip_omml_bytes(docx_path: Path) -> None:
    """Remove any leftover OMML (safety)."""
    buf = BytesIO()
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/document.xml":
                xml = data.decode("utf-8")

                def _text(om: str) -> str:
                    ts = re.findall(r"<m:t[^>]*>(.*?)</m:t>", om, flags=re.S)
                    t = "".join(ts)
                    return (
                        t.replace("&lt;", "<")
                        .replace("&gt;", ">")
                        .replace("&amp;", "&")
                    )

                def repl(m: re.Match) -> str:
                    t = _text(m.group(0))
                    esc = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    sp = ' xml:space="preserve"' if (not t or t[:1].isspace() or t[-1:].isspace()) else ""
                    return (
                        "<w:r><w:rPr>"
                        '<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/>'
                        "<w:i/><w:sz w:val=\"21\"/></w:rPr>"
                        f"<w:t{sp}>{esc}</w:t></w:r>"
                    )

                xml = re.sub(r"<m:oMathPara\b[^>]*>.*?</m:oMathPara>", repl, xml, flags=re.S)
                xml = re.sub(r"<m:oMath\b[^>]*>.*?</m:oMath>", repl, xml, flags=re.S)
                data = xml.encode("utf-8")
            zout.writestr(info, data)
    docx_path.write_bytes(buf.getvalue())


def _run_has_placeholder(run) -> list[int]:
    text = run.text or ""
    return [int(x) for x in re.findall(r"⟦EQ(\d{3})⟧", text)]


def insert_inline_formula_pictures(docx_path: Path, records: list[dict]) -> int:
    """Replace ⟦EQnnn⟧ placeholders with inline pictures."""
    doc = Document(str(docx_path))
    replaced = 0
    by_idx = {r["index"]: r for r in records if not r["display"]}

    def process_paragraph(paragraph) -> None:
        nonlocal replaced
        # Work on a snapshot of runs; rebuilding when placeholder found
        full = "".join(r.text or "" for r in paragraph.runs)
        if "⟦EQ" not in full:
            return
        # Rebuild paragraph runs: split by placeholders
        pattern = re.compile(r"⟦EQ(\d{3})⟧")
        # Clear existing runs
        p_element = paragraph._p
        for child in list(p_element):
            if child.tag == qn("w:r"):
                p_element.remove(child)
        pos = 0
        for m in pattern.finditer(full):
            before = full[pos : m.start()]
            if before:
                run = paragraph.add_run(before)
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                run.font.size = Pt(12)
            idx = int(m.group(1))
            rec = by_idx.get(idx)
            if rec and rec["path"].exists():
                run = paragraph.add_run()
                # inline height ~ body text
                run.add_picture(str(rec["path"]), height=Cm(0.55) if len(rec["latex"]) < 12 else Cm(0.7))
                replaced += 1
            else:
                paragraph.add_run(m.group(0))
            pos = m.end()
        after = full[pos:]
        if after:
            run = paragraph.add_run(after)
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(12)

    for paragraph in doc.paragraphs:
        process_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    process_paragraph(paragraph)

    doc.save(str(docx_path))
    return replaced


def main() -> int:
    if not MD_SRC.exists():
        print(f"ERROR: missing {MD_SRC}", file=sys.stderr)
        return 1
    if not MD2DOCX.exists():
        print(f"ERROR: missing md_to_docx at {MD2DOCX}", file=sys.stderr)
        return 1

    md_text = MD_SRC.read_text(encoding="utf-8")
    md_ready, records = collect_and_render(md_text)
    print(f"rendered {len(records)} formulas → {FORMULA_DIR}")

    with tempfile.TemporaryDirectory(prefix="guide_docx_") as td:
        td_path = Path(td)
        # copy guide_assets so relative image paths resolve
        assets_src = ROOT / "guide_assets"
        assets_dst = td_path / "guide_assets"
        shutil.copytree(assets_src, assets_dst)
        md_tmp = td_path / "构建指南.for_docx.md"
        md_tmp.write_text(md_ready, encoding="utf-8")
        out_tmp = td_path / "out.docx"

        # import and call convert directly to avoid subprocess path issues
        sys.path.insert(0, str(MD2DOCX.parent))
        from md_to_docx import convert  # type: ignore

        convert(md_tmp, out_tmp)
        strip_omml_bytes(out_tmp)
        n = insert_inline_formula_pictures(out_tmp, records)
        print(f"inline formula pictures inserted: {n}")

        # final OMML strip after picture insert (noop usually)
        strip_omml_bytes(out_tmp)

        targets = [
            ROOT / "构建指南.docx",
            ROOT / "构建指南_兼容版.docx",
            ROOT / "构建指南_含公式修复版.docx",
        ]
        for t in targets:
            shutil.copy2(out_tmp, t)
            print(f"saved {t.name} ({t.stat().st_size} bytes)")

    # verify
    with zipfile.ZipFile(ROOT / "构建指南.docx") as z:
        xml = z.read("word/document.xml").decode("utf-8")
    print(f"verify oMath={xml.count('<m:oMath')} placeholders_left={xml.count('⟦EQ')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
